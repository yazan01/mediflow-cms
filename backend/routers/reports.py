# backend/routers/reports.py
from datetime import datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
import io
import csv

from database import get_db
from auth import get_current_user, require_roles
import models

router = APIRouter(prefix="/api/reports", tags=["reports"])

REPORT_ROLES = ("SUPER_ADMIN", "CLINIC_MANAGER", "ACCOUNTANT", "AUDITOR")


def _parse_date(value: Optional[str], field: str) -> Optional[datetime]:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except (ValueError, TypeError):
        return None


def _safe_csv(value: str) -> str:
    """Prefix formula-injection characters to prevent spreadsheet execution."""
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _build_stats(db: Session, start: datetime, end: Optional[datetime], prev_start: datetime, prev_end: datetime, doctor_id: Optional[str] = None, branch_id: Optional[str] = None):
    now = end or datetime.utcnow()

    pay_q = db.query(func.sum(models.Payment.amount)).filter(models.Payment.paidAt >= start)
    if end:
        pay_q = pay_q.filter(models.Payment.paidAt < end)
    revenue_mtd = float(pay_q.scalar() or 0)

    prev_rev = float(
        db.query(func.sum(models.Payment.amount))
        .filter(models.Payment.paidAt >= prev_start, models.Payment.paidAt < prev_end)
        .scalar() or 0
    )

    appt_q = db.query(func.count(models.Appointment.id)).filter(models.Appointment.scheduledAt >= start)
    if end:
        appt_q = appt_q.filter(models.Appointment.scheduledAt < end)
    if doctor_id:
        appt_q = appt_q.filter(models.Appointment.doctorId == doctor_id)
    if branch_id:
        appt_q = appt_q.filter(models.Appointment.branchId == branch_id)
    appointments_mtd = appt_q.scalar() or 0

    prev_appt_q = db.query(func.count(models.Appointment.id)).filter(
        models.Appointment.scheduledAt >= prev_start, models.Appointment.scheduledAt < prev_end
    )
    if doctor_id:
        prev_appt_q = prev_appt_q.filter(models.Appointment.doctorId == doctor_id)
    prev_appointments = prev_appt_q.scalar() or 0

    pat_q = db.query(func.count(models.Patient.id)).filter(models.Patient.createdAt >= start)
    if end:
        pat_q = pat_q.filter(models.Patient.createdAt < end)
    patients_mtd = pat_q.scalar() or 0

    prev_pat = db.query(func.count(models.Patient.id)).filter(
        models.Patient.createdAt >= prev_start, models.Patient.createdAt < prev_end
    ).scalar() or 0

    ns_q = db.query(func.count(models.Appointment.id)).filter(
        models.Appointment.scheduledAt >= start, models.Appointment.status == "NO_SHOW"
    )
    if end:
        ns_q = ns_q.filter(models.Appointment.scheduledAt < end)
    no_show_count = ns_q.scalar() or 0
    no_show_rate = round((no_show_count / appointments_mtd * 100), 1) if appointments_mtd > 0 else 0.0

    prev_ns = db.query(func.count(models.Appointment.id)).filter(
        models.Appointment.scheduledAt >= prev_start, models.Appointment.scheduledAt < prev_end, models.Appointment.status == "NO_SHOW"
    ).scalar() or 0
    prev_ns_rate = round((prev_ns / prev_appointments * 100), 1) if prev_appointments > 0 else 0.0

    def pct(cur, prev):
        if prev == 0:
            return 0
        return round((cur - prev) / prev * 100, 1)

    return {
        "revenueMTD": revenue_mtd,
        "patientsMTD": patients_mtd,
        "appointmentsMTD": appointments_mtd,
        "noShowRate": no_show_rate,
        "revenueChange": pct(revenue_mtd, prev_rev),
        "patientsChange": pct(patients_mtd, prev_pat),
        "appointmentsChange": pct(appointments_mtd, prev_appointments),
        "noShowChange": pct(no_show_rate, prev_ns_rate),
    }


@router.get("/overview")
def get_reports_overview(
    range: str = Query("this_month"),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    doctor_id: Optional[str] = Query(None),
    branch_id: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(require_roles(*REPORT_ROLES)),
):
    now = datetime.utcnow()

    if date_from:
        start = _parse_date(date_from, "date_from") or now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        end = _parse_date(date_to, "date_to")
        period_days = (end - start).days if end else (now - start).days
        prev_start = start - timedelta(days=period_days or 30)
        prev_end = start
    elif range == "last_year":
        start = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        end = None
        prev_start = start.replace(year=start.year - 1)
        prev_end = start
    elif range == "last_quarter":
        start = now - timedelta(days=90)
        end = None
        prev_start = start - timedelta(days=90)
        prev_end = start
    elif range == "last_month":
        first_of_month = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        prev_end = first_of_month
        prev_month = first_of_month - timedelta(days=1)
        prev_start = prev_month.replace(day=1)
        start = prev_start
        end = prev_end
    else:  # this_month
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        end = None
        prev_end = start
        prev_month = start - timedelta(days=1)
        prev_start = prev_month.replace(day=1)

    return _build_stats(db, start, end, prev_start, prev_end, doctor_id, branch_id)


@router.get("/data")
def get_report_data(
    report_type: str = Query("payments"),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    doctor_id: Optional[str] = Query(None),
    branch_id: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    sort_by: Optional[str] = Query(None),
    sort_dir: str = Query("desc"),
    db: Session = Depends(get_db),
    _user=Depends(require_roles(*REPORT_ROLES)),
):
    """Generic paginated/filterable data endpoint for the report viewer."""
    now = datetime.utcnow()
    start = _parse_date(date_from, "date_from") or now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    end = _parse_date(date_to, "date_to") or now

    if report_type == "payments":
        query = (
            db.query(models.Payment, models.Invoice, models.Patient)
            .join(models.Invoice, models.Payment.invoiceId == models.Invoice.id)
            .join(models.Patient, models.Invoice.patientId == models.Patient.id)
            .filter(models.Payment.paidAt >= start, models.Payment.paidAt < end)
        )
        if status and status != "ALL":
            query = query.filter(models.Invoice.status == status)
        if branch_id and branch_id != "ALL":
            query = query.filter(models.Invoice.branchId == branch_id)
        if search:
            query = query.filter(
                models.Invoice.invoiceNo.ilike(f"%{search}%") |
                models.Patient.firstName.ilike(f"%{search}%") |
                models.Patient.lastName.ilike(f"%{search}%")
            )
        total = query.count()
        rows = query.order_by(models.Payment.paidAt.desc()).offset((page - 1) * page_size).limit(page_size).all()
        data = [{
            "invoiceNo": p[1].invoiceNo,
            "patient": f"{p[2].firstName} {p[2].lastName}",
            "mrn": p[2].mrn,
            "amount": float(p[0].amount),
            "method": p[0].method,
            "status": p[1].status,
            "date": p[0].paidAt.isoformat() if p[0].paidAt else None,
        } for p in rows]
        columns = ["invoiceNo", "patient", "mrn", "amount", "method", "status", "date"]

    elif report_type == "appointments":
        query = (
            db.query(models.Appointment)
            .options(
                joinedload(models.Appointment.patient),
                joinedload(models.Appointment.doctor).joinedload(models.Doctor.user),
                joinedload(models.Appointment.branch),
            )
            .filter(models.Appointment.scheduledAt >= start, models.Appointment.scheduledAt < end)
        )
        if doctor_id and doctor_id != "ALL":
            query = query.filter(models.Appointment.doctorId == doctor_id)
        if branch_id and branch_id != "ALL":
            query = query.filter(models.Appointment.branchId == branch_id)
        if status and status != "ALL":
            query = query.filter(models.Appointment.status == status)
        if search:
            query = query.filter(
                models.Appointment.id.ilike(f"%{search}%") |
                models.Patient.firstName.ilike(f"%{search}%") |
                models.Patient.lastName.ilike(f"%{search}%")
            )
        total = query.count()
        rows = query.order_by(models.Appointment.scheduledAt.desc()).offset((page - 1) * page_size).limit(page_size).all()
        data = [{
            "patient": f"{a.patient.firstName} {a.patient.lastName}" if a.patient else "",
            "mrn": a.patient.mrn if a.patient else "",
            "doctor": a.doctor.user.name if a.doctor and a.doctor.user else "",
            "scheduledAt": a.scheduledAt.isoformat() if a.scheduledAt else None,
            "type": a.type,
            "status": a.status,
            "branch": a.branch.name if a.branch else "",
        } for a in rows]
        columns = ["patient", "mrn", "doctor", "scheduledAt", "type", "status", "branch"]

    elif report_type == "patients":
        query = db.query(models.Patient).filter(
            models.Patient.deletedAt == None,  # noqa: E711
            models.Patient.createdAt >= start,
            models.Patient.createdAt < end,
        )
        if search:
            query = query.filter(
                models.Patient.firstName.ilike(f"%{search}%") |
                models.Patient.lastName.ilike(f"%{search}%") |
                models.Patient.mrn.ilike(f"%{search}%")
            )
        if status and status != "ALL":
            query = query.filter(models.Patient.isActive == (status == "ACTIVE"))
        total = query.count()
        rows = query.order_by(models.Patient.createdAt.desc()).offset((page - 1) * page_size).limit(page_size).all()
        data = [{
            "mrn": p.mrn,
            "name": f"{p.firstName} {p.lastName}",
            "gender": p.gender,
            "phone": p.phone,
            "nationality": p.nationality,
            "bloodType": p.bloodType,
            "status": "ACTIVE" if p.isActive else "INACTIVE",
            "createdAt": p.createdAt.isoformat() if p.createdAt else None,
        } for p in rows]
        columns = ["mrn", "name", "gender", "phone", "nationality", "bloodType", "status", "createdAt"]

    else:
        raise HTTPException(status_code=400, detail=f"Unknown report_type: {report_type!r}")

    return {
        "data": data,
        "columns": columns,
        "total": total,
        "page": page,
        "pageSize": page_size,
        "totalPages": -(-total // page_size),
    }


@router.get("/export")
def export_report(
    report_type: str = Query("payments"),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    doctor_id: Optional[str] = Query(None),
    branch_id: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    format: str = Query("csv"),
    lang: str = Query("en"),
    db: Session = Depends(get_db),
    _user=Depends(require_roles(*REPORT_ROLES)),
):
    now = datetime.utcnow()
    start = _parse_date(date_from, "date_from") or now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    end = _parse_date(date_to, "date_to") or now

    # ── Column header labels (EN / AR) ──────────────────────────────────────
    HEADERS = {
        "payments": {
            "en": ["Invoice No", "Patient", "MRN", "Amount", "Method", "Status", "Date"],
            "ar": ["رقم الفاتورة", "المريض", "رقم السجل", "المبلغ", "طريقة الدفع", "الحالة", "التاريخ"],
        },
        "appointments": {
            "en": ["Patient", "MRN", "Doctor", "Scheduled At", "Type", "Status", "Branch"],
            "ar": ["المريض", "رقم السجل", "الطبيب", "الموعد المحدد", "النوع", "الحالة", "الفرع"],
        },
        "patients": {
            "en": ["MRN", "Name", "Gender", "Phone", "Nationality", "Blood Type", "Status", "Created"],
            "ar": ["رقم السجل", "الاسم", "الجنس", "الهاتف", "الجنسية", "فصيلة الدم", "الحالة", "تاريخ التسجيل"],
        },
    }
    headers_dict = HEADERS.get(report_type, HEADERS["payments"])
    headers = headers_dict.get(lang, headers_dict["en"])

    # ── Fetch all matching rows (no pagination for export) ──────────────────
    if report_type == "payments":
        query = (
            db.query(models.Payment, models.Invoice, models.Patient)
            .join(models.Invoice, models.Payment.invoiceId == models.Invoice.id)
            .join(models.Patient, models.Invoice.patientId == models.Patient.id)
            .filter(models.Payment.paidAt >= start, models.Payment.paidAt < end)
        )
        if status and status != "ALL":
            query = query.filter(models.Invoice.status == status)
        if branch_id and branch_id != "ALL":
            query = query.filter(models.Invoice.branchId == branch_id)
        if search:
            query = query.filter(
                models.Invoice.invoiceNo.ilike(f"%{search}%") |
                models.Patient.firstName.ilike(f"%{search}%") |
                models.Patient.lastName.ilike(f"%{search}%")
            )
        rows_raw = query.order_by(models.Payment.paidAt.desc()).all()
        rows = [[
            p[1].invoiceNo or "",
            f"{p[2].firstName} {p[2].lastName}",
            p[2].mrn or "",
            str(float(p[0].amount)),
            p[0].method or "",
            p[1].status or "",
            p[0].paidAt.strftime("%Y-%m-%d %H:%M") if p[0].paidAt else "",
        ] for p in rows_raw]

    elif report_type == "appointments":
        query = (
            db.query(models.Appointment)
            .options(
                joinedload(models.Appointment.patient),
                joinedload(models.Appointment.doctor).joinedload(models.Doctor.user),
                joinedload(models.Appointment.branch),
            )
            .filter(models.Appointment.scheduledAt >= start, models.Appointment.scheduledAt < end)
        )
        if doctor_id and doctor_id != "ALL":
            query = query.filter(models.Appointment.doctorId == doctor_id)
        if branch_id and branch_id != "ALL":
            query = query.filter(models.Appointment.branchId == branch_id)
        if status and status != "ALL":
            query = query.filter(models.Appointment.status == status)
        if search:
            q_search = f"%{search}%"
            query = query.filter(
                models.Patient.firstName.ilike(q_search) |
                models.Patient.lastName.ilike(q_search)
            )
        rows_raw = query.order_by(models.Appointment.scheduledAt.desc()).all()
        rows = [[
            f"{a.patient.firstName} {a.patient.lastName}" if a.patient else "",
            a.patient.mrn if a.patient else "",
            a.doctor.user.name if a.doctor and a.doctor.user else "",
            a.scheduledAt.strftime("%Y-%m-%d %H:%M") if a.scheduledAt else "",
            a.type or "",
            a.status or "",
            a.branch.name if a.branch else "",
        ] for a in rows_raw]

    elif report_type == "patients":
        query = db.query(models.Patient).filter(
            models.Patient.deletedAt == None,  # noqa: E711
            models.Patient.createdAt >= start,
            models.Patient.createdAt < end,
        )
        if search:
            query = query.filter(
                models.Patient.firstName.ilike(f"%{search}%") |
                models.Patient.lastName.ilike(f"%{search}%") |
                models.Patient.mrn.ilike(f"%{search}%")
            )
        if status and status != "ALL":
            query = query.filter(models.Patient.isActive == (status == "ACTIVE"))
        rows_raw = query.order_by(models.Patient.createdAt.desc()).all()
        rows = [[
            p.mrn or "",
            f"{p.firstName} {p.lastName}",
            p.gender or "",
            p.phone or "",
            p.nationality or "",
            p.bloodType or "",
            "ACTIVE" if p.isActive else "INACTIVE",
            p.createdAt.strftime("%Y-%m-%d") if p.createdAt else "",
        ] for p in rows_raw]

    else:
        raise HTTPException(status_code=400, detail=f"Unknown report_type: {report_type!r}")

    fname_base = f"{report_type}_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}"

    # ── CSV export ───────────────────────────────────────────────────────────
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_safe_csv(str(v)) for v in row])
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.csv"'},
        )

    # ── Excel (xlsx) export ─────────────────────────────────────────────────
    if format == "xlsx":
        try:
            import openpyxl
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side, numbers
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="openpyxl not installed on server")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report_type.capitalize()

        # RTL sheet direction for Arabic
        if lang == "ar":
            ws.sheet_view.rightToLeft = True

        # Styles
        header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="002045")
        header_align = Alignment(
            horizontal="right" if lang == "ar" else "left",
            vertical="center",
            wrap_text=False,
        )
        thin = Side(style="thin", color="C4C6CF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        alt_fill = PatternFill("solid", fgColor="F4F3F7")
        normal_align = Alignment(
            horizontal="right" if lang == "ar" else "left",
            vertical="center",
        )

        # Header row
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        # Data rows
        for row_idx, row in enumerate(rows, 2):
            fill = alt_fill if row_idx % 2 == 0 else None
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = normal_align
                cell.border = border
                if fill:
                    cell.fill = fill

        # Auto-size columns (cap at 60 chars)
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    cell_len = len(str(cell.value)) if cell.value else 0
                    if cell_len > max_len:
                        max_len = cell_len
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 64)

        # Freeze header row
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.xlsx"'},
        )

    raise HTTPException(status_code=400, detail=f"Unsupported format: {format!r}. Use 'csv' or 'xlsx'.")
